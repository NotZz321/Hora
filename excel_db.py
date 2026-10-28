"""ใช้ไฟล์ Excel (data/hora.xlsx) เป็น database

โครงสร้างออกแบบให้ต่อ Tableau ได้ทันที
  sheet "charts"  : 1 แถว = 1 ดวง (ข้อมูลหลัก ใช้ทำสถิติ/กราฟ)
  sheet "planets" : 1 แถว = ดาว 1 ดวงของ 1 ดวงชาตา (long format, เชื่อมกับ charts ด้วย chart_id)
"""
import os
import re
import sys
from datetime import datetime, date
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'lib'))
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data', 'hora.xlsx')

CHART_COLS = [
    'chart_id', 'แหล่งข้อมูล', 'บันทึกเมื่อ', 'ชื่อเล่น',
    'วันเกิด', 'ปีเกิด_พศ', 'เดือนเกิด', 'เวลาเกิด', 'ชั่วโมงเกิด', 'จังหวัดเกิด',
    'ลัคนา', 'ชาวราศี', 'ตนุลัคน์', 'ตนุเศษ',
    'ภพกัมมะ_ราศี', 'ดาวกัมมะ', 'ดาวกัมมะ_ตำแหน่ง',
    'ดาวทำนาย_1', 'ดาวทำนาย_2', 'ดาวทำนาย_3',
    'อาชีพทำนาย_1', 'อาชีพทำนาย_2', 'อาชีพทำนาย_3',
]
ID_PREFIX = 'X'          # chart_id แบบ X001, X002, ... (เกิน 999 จะเป็น X1000)
_ID_RE = re.compile(r'^' + ID_PREFIX + r'(\d+)$')
PLANET_COLS = ['chart_id', 'ดาว_เลข', 'ดาว', 'ราศี', 'ราศี_ลำดับ', 'องศาในราศี',
               'ลองจิจูด', 'ภพ', 'ตำแหน่ง']


class ExcelLockedError(Exception):
    """ไฟล์ Excel ถูกเปิดค้างอยู่ในโปรแกรม Excel จึงบันทึกไม่ได้"""


def _style_header(ws, cols):
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='5B3E8F')
        cell.alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'A2'
    for i, name in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(name) + 4)


def _backup_if_old_header():
    """ถ้าไฟล์เดิมมีหัวคอลัมน์ไม่ตรงกับ CHART_COLS (เพราะแก้โค้ด) เปลี่ยนชื่อเป็นไฟล์สำรอง แล้วให้สร้างใหม่"""
    try:
        wb = load_workbook(DB_PATH, read_only=True)
        head = list(next(wb['charts'].iter_rows(max_row=1, values_only=True)))
        wb.close()
    except PermissionError:
        raise ExcelLockedError('ไฟล์ hora.xlsx เปิดค้างอยู่ใน Excel กรุณาปิดไฟล์แล้วลองใหม่')
    if head == CHART_COLS:
        return
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    backup = os.path.join(os.path.dirname(DB_PATH), f'hora_backup_{stamp}.xlsx')
    try:
        os.rename(DB_PATH, backup)
    except PermissionError:
        raise ExcelLockedError('ไฟล์ hora.xlsx เปิดค้างอยู่ใน Excel กรุณาปิดไฟล์แล้วลองใหม่')


def init_db():
    """สร้างไฟล์ Excel พร้อมหัวตาราง ถ้ายังไม่มี (ไฟล์เดิมที่หัวคอลัมน์ไม่ตรงจะถูกสำรองไว้ก่อน)"""
    if os.path.exists(DB_PATH):
        _backup_if_old_header()
    if os.path.exists(DB_PATH):
        return
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    wb = Workbook()
    ws = wb.active; ws.title = 'charts'
    _style_header(ws, CHART_COLS)
    _style_header(wb.create_sheet('planets'), PLANET_COLS)
    wb.save(DB_PATH)


def _save(wb):
    try:
        wb.save(DB_PATH)
    except PermissionError:
        raise ExcelLockedError('ไฟล์ hora.xlsx เปิดค้างอยู่ใน Excel กรุณาปิดไฟล์แล้วลองใหม่')


def next_chart_id(existing):
    """รหัสถัดไปจากรหัสที่มีอยู่ เช่น ['X001', 'X007'] → 'X008' (ไม่ซ้ำแม้มีคนลบแถวใน Excel)"""
    nums = [int(m.group(1)) for v in existing if isinstance(v, str) and (m := _ID_RE.match(v))]
    return f'{ID_PREFIX}{max(nums, default=0) + 1:03d}'


def _open_wb():
    init_db()
    try:
        return load_workbook(DB_PATH)
    except PermissionError:
        raise ExcelLockedError('ไฟล์ hora.xlsx เปิดค้างอยู่ใน Excel กรุณาปิดไฟล์แล้วลองใหม่')


def _append(ws, wp, chart_id, result, nickname, source, saved_at):
    i, k = result['input'], result['karma']
    top = [t['group'] for t in result['top_careers']] + ['', '', '']
    stars = [t['title'] for t in result['top_careers']] + ['', '', '']
    ws.append([
        chart_id, source, saved_at, nickname,
        date(i['year_be'] - 543, i['month'], i['day']), i['year_be'], i['month'],
        f"{i['hh']:02d}:{i['mm']:02d}", i['hh'], i['province'],
        result['lagna']['sign'], result['sun_sign'],
        result['tanu_lagna']['name'], result['tanuset']['name'],
        k['sign'], k['lord_name'], k['dignity'],
        stars[0], stars[1], stars[2],
        top[0], top[1], top[2],
    ])
    ws.cell(ws.max_row, 3).number_format = 'yyyy-mm-dd hh:mm'
    ws.cell(ws.max_row, 5).number_format = 'yyyy-mm-dd'
    for p in result['planets']:
        wp.append([chart_id, p['no'], p['name'], p['sign'], p['sign_no'] + 1,
                   p['deg_in_sign'], p['lon'], p['house'], p['dignity']])


def save_chart(result, nickname='', source='จริง'):
    """บันทึกผลจาก career.analyze() ลง Excel คืน chart_id (เช่น 'X001')"""
    return save_many([(result, nickname, source, datetime.now().replace(microsecond=0))])[0]


def save_many(items):
    """บันทึกหลายดวงในการเปิด-ปิดไฟล์ครั้งเดียว (ใช้กับ mock data)
    items = [(result, nickname, source, saved_at), ...]  คืน list ของ chart_id"""
    wb = _open_wb()
    ws, wp = wb['charts'], wb['planets']
    existing = [r[0] for r in ws.iter_rows(min_row=2, max_col=1, values_only=True)]
    ids = []
    for result, nickname, source, saved_at in items:
        cid = next_chart_id(existing + ids)
        _append(ws, wp, cid, result, nickname, source, saved_at)
        ids.append(cid)
    _save(wb)
    return ids


def count_source(source):
    """นับจำนวนดวงตามแหล่งข้อมูล (เช่น 'mock')"""
    return sum(1 for r in list_charts() if r.get('แหล่งข้อมูล') == source)


def list_charts():
    """อ่าน sheet charts ทั้งหมด คืน list ของ dict (ใหม่สุดก่อน)"""
    init_db()
    wb = load_workbook(DB_PATH, read_only=True)
    rows = list(wb['charts'].iter_rows(values_only=True))
    wb.close()
    head, data = rows[0], rows[1:]
    return [dict(zip(head, r)) for r in reversed(data) if r[0] is not None]


def get_chart(chart_id):
    for row in list_charts():
        if row['chart_id'] == chart_id:
            return row
    return None
