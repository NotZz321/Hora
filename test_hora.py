import os
import sys
import tempfile
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, 'lib'))

import suriyayatra_v3 as S
import career as C

passed = failed = 0


def check(name, got, expected):
    global passed, failed
    if got == expected:
        passed += 1
        print(f'  ✓ {name}')
    else:
        failed += 1
        print(f'  ✗ {name}: ได้ {got!r} ต้องการ {expected!r}')


def sign_of(d, hh, mm, planet, lon=100.5):
    pos, _ = S.chart(d, hh, mm, lon)
    return S.SIGNS[int(pos[planet] // 30)]


print('1) ตัวอย่างในชีทครั้งที่ 10: 20 เม.ย. 2549 23:20 น.')
pos, lg, frame = S.horoscope(20, 4, 2549, 23, 20)
check('ชาวราศีเมษ', S.SIGNS[int(pos[1] // 30)], 'เมษ')
check('ลัคนาธนู', S.SIGNS[int(lg // 30)], 'ธนู')

print('2) สูตรตนุเศษ (ชีทครั้งที่ 10)')


def tanuset_case(A, B):
    lag = 0
    t = S.KASET[lag]
    ts = (lag + A - 1) % 12
    signs = {p: 6 for p in S.PLANET_NAMES}
    signs[t] = ts
    lord = S.KASET[ts]
    if lord != t:
        signs[lord] = (ts + B - 1) % 12
    return S.life_frame(signs, lag)['ตนุเศษ'][0]


check('9×2 เศษ 4 → ๔', tanuset_case(9, 2), 4)
check('9×3 เศษ 6 → ๖', tanuset_case(9, 3), 6)
check('9×4 เศษ 1 → ๑', tanuset_case(9, 4), 1)
check('ตนุลัคน์กุมลัคนา → ๑', tanuset_case(1, 1), 1)
check('ตนุลัคน์เล็งลัคนา → ๗', tanuset_case(7, 2), 7)

print('3) ตำแหน่งมาตรฐานดาว (ชีทครั้งที่ 07, 10, 11)')
SIGN_NO = {n: i for i, n in enumerate(S.SIGNS)}


def names(p, sign_name):
    s = SIGN_NO[sign_name]
    return {n for n, t in C.POSITIONS.items() if s in t.get(p, ())}


check('๑ ในกรกฎ = มหาจักร', 'มหาจักร' in names(1, 'กรกฎ'), True)
check('๕ ในตุลย์ = เทวีโชค', 'เทวีโชค' in names(5, 'ตุลย์'), True)
check('๗ ในตุลย์ = อุจจ์', 'อุจจ์' in names(7, 'ตุลย์'), True)
check('๓ ในตุลย์ = ประเกษตร', 'ประเกษตร' in names(3, 'ตุลย์'), True)
check('๘ ในมีน = ปกติ', names(8, 'มีน'), set())
check('๗ ในกุมภ์ = เกษตร', 'เกษตร' in names(7, 'กุมภ์'), True)
check('๗ ในสิงห์ ไม่ใช่ประเกษตร', 'ประเกษตร' in names(7, 'สิงห์'), False)
check('๘ ในสิงห์ = ประเกษตร', 'ประเกษตร' in names(8, 'สิงห์'), True)
for p, sign, expect in [
        (3, 'พิจิก', {'เกษตร', 'เทวีโชค'}), (4, 'กันย์', {'เกษตร', 'อุจจ์'}),
        (7, 'พิจิก', {'จุลจักร', 'ราชาโชค'}), (7, 'พฤษภ', {'มหาจักร', 'เทวีโชค'}),
        (4, 'มีน', {'ประเกษตร', 'นิจ'}), (6, 'เมษ', {'อุจจาภิมุข', 'ประเกษตร'}),
        (4, 'สิงห์', {'ราชาโชค', 'อุจจาวิลาส'}), (2, 'เมถุน', {'อุจจาภิมุข'}),
        (2, 'ตุลย์', {'จุลจักร'}), (8, 'มังกร', {'มหาจักร'})]:
    check(f'เฉลย: {p} ใน{sign} มี {"/".join(sorted(expect))}', expect <= names(p, sign), True)
check('ทุกตำแหน่งมีครบดาว ๑–๘', all(set(t) == set(range(1, 9)) for t in C.POSITIONS.values()), True)
check('ราชาโชค ๕− = ถูกนินทา', 'ถูกนินทา' in C.dignity(5, 0, k5_good=False), True)
check('เทวีโชค ๕+ = โชคจากเสน่ห์', 'โชคจากเสน่ห์' in C.dignity(5, 6, k5_good=True), True)

print('4) วันย้ายราศีจริง (ใช้ปรับเทียบ ทดสอบ 1 วันหลังย้าย)')
check('ราหูเข้ากุมภ์ 5 พ.ค. 68', sign_of(date(2025, 5, 6), 12, 0, 8), 'กุมภ์')
check('ราหูยังอยู่กุมภ์ 21 พ.ย. 69', sign_of(date(2026, 11, 21), 12, 0, 8), 'กุมภ์')
check('เสาร์เข้ามีน 19 พ.ค. 68', sign_of(date(2025, 5, 20), 12, 0, 7), 'มีน')
check('พฤหัสเข้ากรกฎ 1 มิ.ย. 69', sign_of(date(2026, 6, 2), 12, 0, 5), 'กรกฎ')
check('เกตุเข้าเมถุน 30 ต.ค. 64', sign_of(date(2021, 10, 31), 12, 0, 9), 'เมถุน')
check('เกตุเข้ากรกฎ 14 ก.ค. 66', sign_of(date(2023, 7, 15), 12, 0, 9), 'กรกฎ')
check('มฤตยูเข้าพฤษภ 7 ก.ค. 65', sign_of(date(2022, 7, 9), 12, 0, 0), 'พฤษภ')

print('5) กรอบอาชีพ ดวงตัวอย่าง 20/4/2549')
r = C.analyze(20, 4, 2549, 23, 20)
check('ภพกัมมะ = กันย์', r['karma']['sign'], 'กันย์')
check('ดาวกัมมะ = ๔', r['karma']['lord'], 4)
check('3 อันดับ = ๔ ๕ ๘', [t['planet'] for t in r['top_careers']], [4, 5, 8])

print('5.1) เกณฑ์ 5 (หลักการที่ผู้ใช้กำหนด: ๕ นิจ = เสื่อม, ที่เหลือ = ทรงคุณภาพ)')
check('๕ ในมังกร (นิจ) = เสื่อมคุณภาพ', C.kesana5({5: 9})['good'], False)
check('๕ ในกรกฎ (อุจจ์) = ทรงคุณภาพ', C.kesana5({5: 3})['good'], True)
check('๕ ในตุลย์ = ทรงคุณภาพ', C.kesana5({5: 6})['good'], True)
check('analyze() มีเกณฑ์ 5', 'kesana5' in r, True)

print('6) Excel บันทึกแล้วอ่านกลับได้ (ใช้ไฟล์ชั่วคราว ไม่แตะ data/hora.xlsx)')
import excel_db as DB
DB.DB_PATH = os.path.join(tempfile.mkdtemp(), 'test.xlsx')
cid = DB.save_chart(r, 'ทดสอบ')
row = DB.get_chart(cid)
check('chart_id แรก = X001', cid, 'X001')
check('chart_id ถัดไป = X002', DB.save_chart(r), 'X002')
check('รหัสถัดไปไม่ซ้ำแม้ลบแถว', DB.next_chart_id(['X001', 'X007', None, 'ข้อความ']), 'X008')
check('ลัคนาใน Excel', row['ลัคนา'], 'ธนู')
for gone in ['ลัคนา_องศา', 'ตนุลัคน์_ภพ', 'ตนุเศษ_ภพ', 'ดาวกัมมะ_ภพ', 'ดาวกัมมะ_คุณภาพ',
             'สายงานจริง', 'ทำนายตรงอันดับ1', 'ทำนายตรงใน3อันดับ']:
    check(f'ไม่เก็บคอลัมน์ {gone}', gone in DB.CHART_COLS, False)
from openpyxl import load_workbook as _lw
check('sheet planets ใช้รหัส X001', next(_lw(DB.DB_PATH)['planets'].iter_rows(min_row=2, values_only=True))[0], 'X001')

check('ข้อมูลจากเว็บ แหล่งข้อมูล = จริง', row['แหล่งข้อมูล'], 'จริง')

print('6.1) mock data (make_mock.py)')
import make_mock as M
m1, m2 = M.build(20), M.build(20)
check('seed เดิม → ชุดเดิม', [x[0]['input'] for x in m1], [x[0]['input'] for x in m2])
check('จังหวัดใน mock ตรงกับรายชื่อจังหวัด', set(M.POPULATION), set(S.PROVINCES))
check('ปีเกิดอยู่ในช่วง', all(M.YEAR_BE[0] <= x[0]['input']['year_be'] <= M.YEAR_BE[1] for x in m1), True)
ids = DB.save_many(m1)
check('mock ต่อรหัสจากข้อมูลจริง', ids[0], 'X003')
check('นับ mock ได้', DB.count_source('mock'), 20)
check('--replace ลบ mock เดิม', M.remove_mock(), 20)
check('หลังลบ mock เหลือแต่ข้อมูลจริง', [x['แหล่งข้อมูล'] for x in DB.list_charts()], ['จริง', 'จริง'])
check('planets ของ mock ถูกลบด้วย', sum(1 for r in _lw(DB.DB_PATH)['planets'].iter_rows(min_row=2)), 20)

print('7) กลุ่มดาวและคำบรรยาย (ชีทครั้งที่ 03)')
check('ชื่อกลุ่มครบ 10 ดวง', set(C.PLANET_TITLE), set(range(10)))
check('คำบรรยายครบ 10 ดวง', all(C.PLANET_PROFILE[p]['desc'] and C.PLANET_PROFILE[p]['good']
                                 and C.PLANET_PROFILE[p]['bad'] for p in range(10)), True)
check('คำบรรยายขึ้นต้นด้วยชื่อกลุ่ม', all(C.PLANET_PROFILE[p]['desc'].startswith(C.PLANET_TITLE[p])
                                     for p in range(10)), True)
check('อันดับ 1 ของดวงตัวอย่าง = พระพุธ', r['main']['title'], 'พระพุธ')
check('อันดับ 2–3 มีชื่อดาว', [t['title'] for t in r['top_careers'][1:]], ['พระพฤหัสบดี', 'พระราหู'])
check('Excel เก็บชื่อดาวทำนาย', row['ดาวทำนาย_1'], 'พระพุธ')

print('8) Excel ไฟล์เก่าหัวคอลัมน์ไม่ตรง → สำรองแล้วสร้างใหม่')
from openpyxl import Workbook
old_dir = tempfile.mkdtemp()
DB.DB_PATH = os.path.join(old_dir, 'hora.xlsx')
wb = Workbook(); wb.active.title = 'charts'; wb.active.append(['chart_id', 'หัวเก่า']); wb.create_sheet('planets')
wb.save(DB.DB_PATH)
DB.save_chart(r)
check('มีไฟล์สำรอง', any(f.startswith('hora_backup_') for f in os.listdir(old_dir)), True)
check('ไฟล์ใหม่ใช้หัวคอลัมน์ล่าสุด', list(DB.list_charts()[0].keys()), DB.CHART_COLS)

print('9) หน้าเว็บ (ไม่มีหน้าประวัติ ผลแสดงทันที)')
import app as A
client = A.app.test_client()
check('ไม่มี /history', client.get('/history').status_code, 404)
check('ไม่มี /chart/1', client.get('/chart/1').status_code, 404)
page = client.post('/predict', data=dict(day=20, month=4, year_be=2549, hh=23, mm=20,
                                         province='กรุงเทพมหานคร')).get_data(as_text=True)
check('หน้าผลแสดงกลุ่มดาว', 'พระพุธ' in page and 'กลุ่มดาวของคุณ' in page, True)
home = client.get('/').get_data(as_text=True)
check('ฟอร์มใช้คำว่า นาฬิกา', 'นาฬิกา' in home, True)
check('ฟอร์มไม่มีช่องสายงานจริง', 'name="actual"' in home, False)
check('คำโปรยหน้าแรก', 'ให้ดวงชาตา<br>พยากรณ์<span>อาชีพของคุณ</span>' in home, True)
check('มีไอคอนบนแท็บ', 'favicon.svg' in home and client.get('/static/favicon.svg').status_code == 200 and client.get('/static/favicon.png').status_code == 200, True)
check('header ติดบนสุด', "classList.toggle('scrolled'" in home and 'position: sticky' in open(os.path.join(HERE, 'static', 'style.css'), encoding='utf-8').read(), True)
check('header มี HORA ARCHEEP', 'H O R A&nbsp;&nbsp;&nbsp;A R C H E E P' in home, True)
check('ไม่มีคำว่า ดวงที่บันทึก ใน JS', 'ดวงที่บันทึก' in open(os.path.join(HERE, 'static', 'sky.js'), encoding='utf-8').read(), False)

print('10) หน้าตาใหม่จากเพื่อน (รูปเทพ) ยังคงข้อตกลงเดิม')
from deities import DEITIES
check('มีรูปเทพครบดาว ๐–๙', all(os.path.exists(os.path.join(HERE, 'image', DEITIES[p]['image']))
                                for p in range(10)), True)
from urllib.parse import quote
check('เปิดรูปมฤตยูได้ (ชื่อไฟล์มีเว้นวรรค)', client.get('/deities/' + quote(DEITIES[0]['image'])).status_code, 200)
check('เปิดรูปเทพได้', client.get('/deities/' + DEITIES[4]['image']).status_code, 200)
check('ไฟล์อื่นนอก DEITIES เปิดไม่ได้', client.get('/deities/app.py').status_code, 404)
check('ไม่มีลิงก์ดวงที่บันทึก', 'ดวงที่บันทึก' in page + client.get('/').get_data(as_text=True), False)
check('สีหน้าผลตามดาวอันดับ 1', 'data-deity="4"' in page, True)
check('อันดับ 2–3 ไม่มีคะแนนน้ำหนักบนการ์ด', 'career-weight' in page, False)
check('ไม่มีส่วนรวมคะแนน', 'รวมคะแนน' in page, False)
check('อันดับ 2–3 มีตัวอย่างอาชีพ', all(w in page for t in r['top_careers'][1:] for w in t['examples'].split()), True)
check('อันดับ 2–3 เป็นการ์ด', page.count('class="runner-card"'), len(r['top_careers']) - 1)
check('การ์ดมีเลขดาวไทย', all(f'<span class="rc-seal" aria-hidden="true">{"๐๑๒๓๔๕๖๗๘๙"[t["planet"]]}</span>' in page for t in r['top_careers'][1:]), True)
from deities import DEITIES as _D
check('ใต้ชื่อดาวหลักมีเส้นทาง (title)', f'<p class="reading-path">{_D[r["main"]["planet"]]["title"]}</p>' in page, True)
check('หลังการ์ดมี title keyword intro', all(_D[t['planet']][k] in page for t in r['top_careers'][1:] for k in ('title', 'keyword', 'intro')), True)
back = page.split('rc-back')[1].split('</article>')[0]
check('หลังการ์ดไม่มีรูปเทพ', '<img' in back, False)
check('การ์ดกดพลิกได้ด้วยคีย์บอร์ด', 'tabindex="0"' in page and "classList.toggle('flipped')" in page, True)
check('คำบรรยายอันดับ 1 มี 3 ย่อหน้า', all(C.PLANET_PROFILE[p]['desc'].count('\n') == 2 for p in range(10)), True)
import random
random.seed(7)
samples = [C.analyze(random.randint(1, 28), random.randint(1, 12), random.randint(2450, 2580),
                     random.randint(0, 23), random.randint(0, 59)) for _ in range(300)]
check('อันดับ 1 = ดาวกัมมะทุกดวง (สุ่ม 300 ดวง)', all(x['main']['planet'] == x['karma']['lord'] for x in samples), True)
r0 = next(x for x in samples if any(t['planet'] == 0 for t in x['top_careers']))
with A.app.test_request_context():
    from flask import render_template
    html0 = render_template('result.html', r=r0, svg='', nickname='', saved=True, warning=None)
check('มฤตยูในอันดับ 2–3 ใช้รูปเทพ ไม่ใช่ ♅', 'Celestial' in html0 and '♅' not in html0, True)

dash = client.get('/dashboard')
dash_html = dash.get_data(as_text=True)
check('มีหน้า /dashboard', dash.status_code, 200)
check('หน้าแรกมีลิงก์ไปหน้าสถิติ', '/dashboard' in home, True)
import json as _json
_payload = _json.loads(dash_html.split('<script id="hora-data" type="application/json">')[1].split('</script>')[0])
check('dashboard ได้ข้อมูลทุกดวงใน Excel', len(_payload['rows']), len(DB.list_charts()))
check('dashboard แปลงจังหวัดเป็นรหัส ISO', all(r['p'].startswith('TH-') for r in _payload['rows']), True)
check('dashboard มีชื่อจังหวัดไทยครบ 77', len(_payload['th']), 77)
check('ไฟล์แผนที่และกราฟเปิดได้', client.get('/static/geo-th.js').status_code == 200 and client.get('/static/dashboard.js').status_code == 200, True)

print(f'\nผ่าน {passed} / {passed + failed}')
sys.exit(1 if failed else 0)
